# -*- coding: utf-8 -*-
"""
Destinations Azerbaijan Transfers PVT 2026 Excel -> transfers cədvəli (geniş format).

Excel layout (sətir 1 başlıqlar, sətir 4 nəqliyyat, məlumat sətir 6+):
  A=Duty Code, B=A (məs. Airport/Hotel/Baku), C=B (məs. Hotel/Lerik),
  D=Service, E=Distance, F=Start Time, G=Duration, H=Day Schedule, I+=qiymətlər
"""
import openpyxl
import os
import sys
import re

BASE = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(os.path.dirname(BASE), "Destinations Azerbaijan Transfers PVT 2026_31.12.2025.xlsx")
OUTPUT_DIR = os.path.join(BASE, "sql_import_transfers")
BATCH_SIZE = 500

VEHICLE_MAP = [
    (r"sedan.*4.*local speaking|sedan.*local", "sedan_4_local"),
    (r"sedan.*4.*english|sedan.*english.*non fluent", "sedan_4_english"),
    (r"vito.*8.*local", "vito_8_local"),
    (r"vito.*8.*english|vito.*english.*non fluent", "vito_8_english"),
    (r"sprinter.*18", "sprinter_18"),
    (r"isuzu.*27", "isuzu_27"),
    (r"403.*45|mercedes 403", "merc_403_45"),
    (r"trivago.*54", "trivago_54"),
    (r"neoplan.*48", "neoplan_48"),
    (r"guides\s*\(a\)", "guides_a"),
    (r"guides\s*\(o\)", "guides_o"),
]

OUTPUT_COLUMNS = [
    "sedan_4_local", "sedan_4_english", "vito_8_local", "vito_8_english",
    "sprinter_18", "isuzu_27", "merc_403_45", "trivago_54", "neoplan_48",
    "guides_a", "guides_o",
]

META_KEYS = ("duty", "a", "b", "service", "distance", "start_time", "duration", "day_schedule")


def map_vehicle_to_column(vehicle_name):
    v = str(vehicle_name).lower()
    for pattern, col in VEHICLE_MAP:
        if re.search(pattern, v):
            return col
    return None


def escape_sql(s):
    if s is None:
        return None
    return str(s).replace("\\", "\\\\").replace("'", "''").replace("\r", " ").replace("\n", " \\n ").strip()


def to_sql_val(v):
    if v is None or v == "" or (isinstance(v, float) and str(v) == "nan"):
        return "NULL"
    if isinstance(v, (int, float)):
        try:
            if isinstance(v, float) and v == int(v):
                return str(int(v))
            return str(v)
        except (ValueError, TypeError):
            return "NULL"
    return "'" + escape_sql(str(v)) + "'"


def _cell_str(val):
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    s = str(val).replace("\t", " ").replace("\r", " ")
    return " ".join(s.split()).strip()


def _normalize_service(service):
    """Excel-də Private Transfer / PrivateTransfer və s. vahid formata."""
    s = _cell_str(service)
    if not s:
        return "Transfer"
    low = s.lower().replace(" ", "")
    if low in ("privatetransfer", "privatetour", "privatetransfer\t"):
        return "Private Transfer"
    if low == "1waytransfer":
        return "1 way Transfer"
    if low == "2waystransfer":
        return "2 ways Transfer"
    return s


def _parse_duty(val):
    if val is None or _cell_str(val) == "":
        return None
    try:
        n = int(float(val))
        return n if n > 0 else None
    except (ValueError, TypeError):
        return None


def _parse_optional_number(val):
    if val is None or _cell_str(val) == "":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _parse_optional_int(val):
    n = _parse_optional_number(val)
    if n is None:
        return None
    try:
        return int(n)
    except (ValueError, TypeError):
        return None


def _header_matches_layout(ws, row=1):
    labels = []
    for col in range(1, 10):
        labels.append(_cell_str(ws.cell(row, col).value).lower())
    if not labels:
        return False
    text = " ".join(labels)
    has_duty = "duty" in labels[0] or "daily" in labels[0] or "code" in labels[0]
    has_service = any("service" in x for x in labels)
    # Standart: A, B, Service sütunları (Distance/Schedule ola bilər)
    if has_duty and has_service and len(labels) > 2 and labels[1] == "a" and labels[2] == "b":
        return True
    # Qısa layout: Duty, şəhər, təyinat, Service, Day Schedule (E sütununda)
    if has_duty and has_service and len(labels) > 4:
        if "schedule" in labels[4] or "schedule" in text:
            return True
    return False


def _detect_meta_columns(ws, header_row=1):
    """
    Başlıq sətirinə görə meta sütun indeksləri (0-based).
    Standart: Duty,A,B,Service,Distance,Start,Duration,Schedule
    Qısa:    Duty,City/Region,Destination,Service,DaySchedule
    """
    labels = []
    for col in range(1, 13):
        labels.append(_cell_str(ws.cell(header_row, col).value).lower())

    if len(labels) > 2 and labels[1] == "a" and labels[2] == "b":
        return {
            "duty": 0, "a": 1, "b": 2, "service": 3,
            "distance": 4, "start_time": 5, "duration": 6, "day_schedule": 7,
        }

    sched_col = next((i for i, lb in enumerate(labels) if "schedule" in lb), 4)
    return {
        "duty": 0, "a": 1, "b": 2, "service": 3,
        "distance": None, "start_time": None, "duration": None,
        "day_schedule": sched_col,
    }


def _find_vehicle_header_row(ws, max_row=8):
    best_row, best_count = 4, 0
    for r in range(1, max_row + 1):
        count = 0
        for col in range(1, min(ws.max_column + 1, 40)):
            val = ws.cell(r, col).value
            if val and map_vehicle_to_column(str(val).strip()):
                count += 1
        if count > best_count:
            best_count = count
            best_row = r
    return best_row if best_count >= 2 else 4


def _get_vehicle_columns(ws, header_row):
    vehicle_cols = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(header_row, col).value
        if val and str(val).strip():
            out_col = map_vehicle_to_column(str(val).strip())
            if out_col:
                vehicle_cols.append((col - 1, out_col))
    return vehicle_cols


def _find_data_start_row(ws, vehicle_header_row):
    for r in range(vehicle_header_row + 1, min(ws.max_row + 1, vehicle_header_row + 6)):
        a = _cell_str(ws.cell(r, 2).value)
        b = _cell_str(ws.cell(r, 3).value)
        svc = _cell_str(ws.cell(r, 4).value)
        duty = _parse_duty(ws.cell(r, 1).value)
        if duty or a or b or svc:
            return r
    return vehicle_header_row + 2


def pick_sheets_to_import(sheetnames, path=""):
    """Import every valid sheet, including sheets newly added by the admin."""
    return list(sheetnames)


def _row_has_price(row, vehicle_cols):
    for col_idx, _ in vehicle_cols:
        if col_idx < len(row):
            p = _parse_optional_number(row[col_idx])
            if p is not None and p >= 0:
                return True
    return False


def _row_get(row, idx):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _map_meta_row(row, cols):
    """Excel sütunları -> transfers sahələri."""
    duty = _parse_duty(_row_get(row, cols["duty"]))
    point_a = _cell_str(_row_get(row, cols["a"])) or None
    point_b = _cell_str(_row_get(row, cols["b"])) or None
    service = _normalize_service(_row_get(row, cols["service"]))
    distance = _row_get(row, cols.get("distance"))
    start_time = _row_get(row, cols.get("start_time"))
    duration = _row_get(row, cols.get("duration"))
    day_schedule = _row_get(row, cols.get("day_schedule"))

    if distance is not None and _cell_str(distance) == "":
        distance = None
    if start_time is not None and _cell_str(start_time) == "":
        start_time = None
    if duration is not None and _cell_str(duration) == "":
        duration = None
    if day_schedule is not None and _cell_str(day_schedule) == "":
        day_schedule = None

    return {
        "duty_code": duty,
        "a": point_a,
        "b": point_b,
        "service": service,
        "distance": distance,
        "start_time": start_time,
        "duration": duration,
        "day_schedule": day_schedule,
    }


def _row_is_valid(meta):
    return bool(meta["a"] or meta["b"] or meta["service"])


def collect_wide_rows(path, skipped=None):
    """Duty kodu transferin unikal açarıdır (A|B|Service təkrarlana bilər).

    Duty Code boş olan sətirlər atlanır; eyni kod bir neçə vərəqdə varsa ilk vərəq üstündür.
    skipped dict verilərsə atlanan sətirlərin sayı yazılır.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Excel tapılmadı: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    all_rows = []
    seen_duties = set()
    if skipped is None:
        skipped = {}
    skipped.update({"no_duty": 0, "duplicate_duty": 0})
    sheets = pick_sheets_to_import(wb.sheetnames, path)

    for sheet_name in sheets:
        ws = wb[sheet_name]
        if not _header_matches_layout(ws):
            continue

        meta_cols = _detect_meta_columns(ws)
        vehicle_header_row = _find_vehicle_header_row(ws)
        vehicle_cols = _get_vehicle_columns(ws, vehicle_header_row)
        if not vehicle_cols:
            continue

        data_start = _find_data_start_row(ws, vehicle_header_row)
        rows = list(ws.iter_rows(min_row=data_start, values_only=True))

        for excel_row_num, row in enumerate(rows, start=data_start):
            if not row:
                continue
            meta = _map_meta_row(row, meta_cols)
            if not _row_is_valid(meta):
                continue

            duty = meta["duty_code"]

            prices = {col: None for col in OUTPUT_COLUMNS}
            for col_idx, out_col in vehicle_cols:
                if col_idx < len(row):
                    p = _parse_optional_number(row[col_idx])
                    if p is not None and p >= 0:
                        prices[out_col] = p

            if not _row_has_price(row, vehicle_cols):
                continue
            if duty is None:
                skipped["no_duty"] += 1
                continue
            if duty in seen_duties:
                skipped["duplicate_duty"] += 1
                continue
            seen_duties.add(duty)

            all_rows.append({
                "duty_code": duty,
                "a": meta["a"] or "Hotel",
                "b": meta["b"] or "Transfer",
                "service": meta["service"] or "Transfer",
                "distance": meta["distance"],
                "start_time": meta["start_time"],
                "duration": meta["duration"],
                "day_schedule": meta["day_schedule"],
                "prices": prices,
            })

    wb.close()
    return all_rows


def main():
    excel_path = sys.argv[1] if len(sys.argv) > 1 else EXCEL_PATH
    output_dir = sys.argv[2] if len(sys.argv) > 2 else OUTPUT_DIR

    rows = collect_wide_rows(excel_path)
    os.makedirs(output_dir, exist_ok=True)

    cols = ["duty_code", "a", "b", "service", "distance", "start_time", "duration_mins", "day_schedule"] + OUTPUT_COLUMNS
    col_str = ", ".join(cols)

    for start in range(0, len(rows), BATCH_SIZE):
        batch = rows[start:start + BATCH_SIZE]
        part = start // BATCH_SIZE
        out_path = os.path.join(output_dir, f"insert_{part}.sql")

        with open(out_path, "w", encoding="utf-8") as f:
            f.write("-- Destinations Azerbaijan Transfers PVT 2026 -> transfers (wide format)\n")
            f.write("SET NAMES utf8mb4;\n\n")
            for r in batch:
                dist = _parse_optional_number(r["distance"])
                dur = _parse_optional_int(r["duration"])
                start_t = r.get("start_time")
                if start_t is not None:
                    start_t = _cell_str(start_t) or None

                vals = [
                    str(r["duty_code"]),
                    to_sql_val(r["a"]),
                    to_sql_val(r["b"]),
                    to_sql_val(r["service"]),
                    str(dist) if dist is not None else "NULL",
                    to_sql_val(start_t),
                    str(dur) if dur is not None else "NULL",
                    to_sql_val(r["day_schedule"]),
                ]
                for col in OUTPUT_COLUMNS:
                    p = r["prices"].get(col)
                    vals.append(str(p) if p is not None else "NULL")
                val_str = ", ".join(vals)
                f.write(f"INSERT INTO transfers ({col_str}) VALUES ({val_str});\n")

        print(f"Yaradıldı: {out_path} ({len(batch)} sətir)")

    print(f"\nCəmi: {len(rows)} sətir -> {(len(rows) + BATCH_SIZE - 1) // BATCH_SIZE} fayl")


if __name__ == "__main__":
    main()
