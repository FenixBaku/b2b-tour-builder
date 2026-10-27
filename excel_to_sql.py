# -*- coding: utf-8 -*-
import openpyxl
import os
import sys

DEFAULT_EXCEL = os.path.join(os.path.dirname(__file__), "Destinations Azerbaijan Activities 2026_updated 31.12.2025.xlsx")
DEFAULT_OUTPUT = os.path.join(os.path.dirname(__file__), "tour_azerbaijan_activities_2026.sql")
EXCEL_PATH = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL
OUTPUT_SQL = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT


def escape_sql(s):
    """SQL string üçün tək dırnağı escape et"""
    if s is None:
        return None
    return str(s).replace("\\", "\\\\").replace("'", "''").replace("\r", " ").replace("\n", " ")


def to_sql_val(v):
    """Dəyəri SQL literal kimi qaytar"""
    if v is None or v == "":
        return "NULL"
    if isinstance(v, (int, float)):
        try:
            return str(int(v)) if isinstance(v, float) and v == int(v) else str(v)
        except (ValueError, TypeError):
            return "NULL"
    return "'" + escape_sql(v) + "'"


def normalize_city(v):
    """City dəyərini kiçik hərflə, DB formatında qaytar"""
    if v is None or str(v).strip() == "":
        return "baku"
    c = str(v).strip().lower()
    m = {"shamakhi": "shamakhi", "shamaxı": "shamakhi", "qabala": "gabala", "qəbələ": "gabala"}
    return m.get(c, c)


def parse_price_val(v):
    """Qiymət hüceyrəsini float-a çevir (0 daxil)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def detect_price_columns(ws, max_scan_row=6, max_col=30):
    """Başlıq sətirlərindən adult/child/season sütunları (0-based)."""
    col = {"adult": 8, "child": 9, "season": None}
    for r in range(1, max_scan_row + 1):
        for c in range(1, max_col + 1):
            val = ws.cell(r, c).value
            if val is None:
                continue
            label = str(val).strip().lower()
            if not label:
                continue
            if "child" in label or "6-12" in label or "6–12" in label or "6_12" in label:
                col["child"] = c - 1
            elif "adult" in label:
                col["adult"] = c - 1
            elif "season" in label and "price" not in label and "peak" not in label:
                col["season"] = c - 1
    if col["season"] is not None and col["season"] in (col["adult"], col["child"]):
        col["season"] = None
    return col


def map_activity_row(row, fallback_city, col):
    """Excel sətirini DB sahələrinə çevir."""
    name = (row[0] if len(row) > 0 else None) or ""
    name = str(name).strip()
    city_val = row[1] if len(row) > 1 else None
    city = normalize_city(city_val) if (city_val is not None and str(city_val).strip()) else fallback_city
    service = str(row[2] if len(row) > 2 else "").strip() or name
    desc = row[3] if len(row) > 3 else None
    open_t = row[4] if len(row) > 4 else None
    close_t = row[5] if len(row) > 5 else None
    dur = row[6] if len(row) > 6 else None
    if dur is not None and isinstance(dur, (int, float)):
        try:
            dur = int(dur)
        except (ValueError, TypeError):
            dur = None
    slots = row[7] if len(row) > 7 else None

    adult = parse_price_val(row[col["adult"]] if len(row) > col["adult"] else None)
    child = parse_price_val(row[col["child"]] if len(row) > col["child"] else None)
    if child is None and adult is not None:
        child = adult

    season = None
    if col["season"] is not None and len(row) > col["season"]:
        sv = row[col["season"]]
        if sv is not None and str(sv).strip() and not isinstance(sv, (int, float)):
            season = str(sv).strip()

    return (name, city, service, desc, open_t, close_t, dur, slots, season, adult, child)


def sheet_to_city(sheet_name):
    """Vərəq adından city adı çıxar (məs: Activity Baku -> baku)"""
    if not sheet_name:
        return "baku"
    s = str(sheet_name).lower()
    if "baku" in s: return "baku"
    if "shamakhi" in s or "shamaxı" in s or "şamaxı" in s: return "shamakhi"
    if "gabala" in s or "qabala" in s or "qəbələ" in s: return "gabala"
    if "shahdag" in s or "şahdağ" in s: return "shahdag"
    if "sheki" in s or "şəki" in s or "shaki" in s: return "sheki"
    return "baku"


def collect_all_rows(path):
    """Bütün vərəqlərdən məlumat sətirlərini topla (city, adult/child düzgün sütunlarla)."""
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Excel tapılmadı: {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        fallback_city = sheet_to_city(sheet_name)
        col = detect_price_columns(ws)
        for row in rows[4:]:  # məlumat 5-ci sətirdən
            if not row:
                continue
            mapped = map_activity_row(row, fallback_city, col)
            if not mapped[0]:
                continue
            out.append(mapped)
    wb.close()
    return out


def main():
    rows = collect_all_rows(EXCEL_PATH)
    d = os.path.dirname(OUTPUT_SQL)
    if d:
        os.makedirs(d, exist_ok=True)
    with open(OUTPUT_SQL, "w", encoding="utf-8") as f:
        f.write("-- Destinations Azerbaijan Activities 2026 - Excel-dan generasiya edilib\n")
        f.write("-- Hədəf cədvəl: tour_azerbaijan_activities_2026\n")
        f.write("-- Fayl: Destinations Azerbaijan Activities 2026_updated 31.12.2025.xlsx\n")
        f.write("-- city sütunu Excel-dan əlavə edilib\n\n")
        f.write("SET NAMES utf8mb4;\n\n")
        f.write("INSERT INTO `tour_azerbaijan_activities_2026` ")
        f.write("(`name`,`city`,`service`,`description`,`open_time`,`close_time`,")
        f.write("`duration_mins`,`slots`,`season`,`adult_price`,`child_price_6_12`)")
        f.write(" VALUES\n")
        vals = []
        for row in rows:
            name, city, service, desc, open_t, close_t, dur, slots, season, adult, child = row
            slots_str = str(slots).strip() if slots is not None and str(slots).strip() else None
            adult_sql = str(adult) if adult is not None else "NULL"
            child_sql = str(child) if child is not None else "NULL"
            one = "(" + to_sql_val(name) + "," + to_sql_val(city) + "," + to_sql_val(service) + "," + to_sql_val(desc) + "," + to_sql_val(open_t) + "," + to_sql_val(close_t) + "," + (str(dur) if dur is not None else "NULL") + "," + to_sql_val(slots_str) + "," + to_sql_val(season) + "," + adult_sql + "," + child_sql + ")"
            vals.append(one)
        f.write(",\n".join(vals))
        f.write(";\n")
    print(f"Yazıldı: {OUTPUT_SQL} ({len(rows)} sətir)")


if __name__ == "__main__":
    main()
