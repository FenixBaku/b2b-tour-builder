# -*- coding: utf-8 -*-
"""Meals_Separated.xlsx -> meals.sql (price ilə)"""
import openpyxl
import os
import re

EXCEL_PATH = r"c:\Users\user\Downloads\Meals_Separated.xlsx"
OUT_SQL = os.path.join(os.path.dirname(__file__), "migx", "meals.sql")
PRICES_MARKER = "-- Sezon fiyatları (2026) - tour_meal_prices"
# Yazılan şəhər adları sonda kodlara çevrilir (BAK, QBL, SHD, SHK) — lib/cities.php ilə eyni
CITY_CODES_SQL = (
    "\n-- Şəhər adlarını kodlara çevir (lib/cities.php, migrations/021_city_codes.sql)\n"
    "UPDATE `tour_meal_options` SET `city` = CASE LOWER(`city`) WHEN 'baku' THEN 'BAK' WHEN 'qabala' THEN 'QBL' "
    "WHEN 'gabala' THEN 'QBL' WHEN 'shahdag' THEN 'SHD' WHEN 'sheki' THEN 'SHK' ELSE `city` END;\n"
)

CITY_MAP = {
    "qabala": "qabala", "qebele": "qabala", "gabala": "qabala",
    "shahdag": "shahdag", "shahdagh": "shahdag",
    "sheki": "sheki", "baku": "baku",
}


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s).strip())


def norm_name(name):
    """VEG  (4 Main Course) -> VEG (4 Main Course)"""
    return norm(name).replace("  ", " ")


def esc(s):
    return s.replace("\\", "\\\\").replace("'", "''")


def extract_prices_from_excel():
    """Excel-dən (city, meal_type, name) -> (price, price2) xəritəsi"""
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    price_map = {}

    for sn in wb.sheetnames:
        city = CITY_MAP.get(sn.lower().replace(" ", ""), CITY_MAP.get(sn.lower(), "baku"))
        if "shahdag" in sn.lower():
            city = "shahdag"
        if "sheki" in sn.lower():
            city = "sheki"
        if "baku" in sn.lower():
            city = "baku"
        if "qabal" in sn.lower() or "qebele" in sn.lower():
            city = "qabala"

        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Find price columns: look for 'Adult' or numeric in first 5 rows
        price_col1, price_col2 = None, None
        data_start = 0
        for row_idx in range(min(5, len(rows))):
            row = rows[row_idx]
            for c in range(len(row)):
                v = row[c]
                if v and str(v).strip().lower() in ("adult", "child (6-12)"):
                    if price_col1 is None:
                        price_col1 = c
                    elif price_col2 is None and c > price_col1:
                        price_col2 = c
                        break
            # Find first row with numeric price
            for c in range(1, len(row)):
                v = row[c]
                if isinstance(v, (int, float)) and 0 <= v <= 500:
                    if price_col1 is None:
                        price_col1 = c
                    if price_col2 is None and c != price_col1:
                        price_col2 = c
                    break
            if price_col2 is not None or price_col1 is not None:
                data_start = row_idx  # include this row (first data row)
                break

        if price_col1 is None:
            continue
        if price_col2 is None:
            price_col2 = price_col1

        name_col, service_col = 0, 1
        last_name = ""

        for row in rows[data_start:]:
            if not row:
                continue
            name_val = norm(row[name_col]) if name_col < len(row) else ""
            service = norm_name(row[service_col]) if service_col < len(row) else ""
            if name_val:
                last_name = name_val
            if not service:
                continue
            ln = last_name.lower()
            if "lunch" in ln and "dinner" not in ln:
                meal_type = "lunch"
            elif "dinner" in ln or "candle" in ln or "gala" in ln:
                meal_type = "dinner"
            elif "hi tea" in ln or "hitea" in ln:
                meal_type = "lunch"
            else:
                continue

            p1 = row[price_col1] if price_col1 < len(row) else None
            p2 = row[price_col2] if price_col2 < len(row) else None
            if p1 is not None and isinstance(p1, (int, float)):
                price = float(p1)
                price2 = float(p2) if (p2 is not None and isinstance(p2, (int, float))) else price
                key = (city, meal_type, service)
                price_map[key] = (price, price2)

    wb.close()
    return price_map


def get_price(price_map, city, meal_type, name, fallback_city="qabala"):
    """Fiyatı tap; shahdag üçün qabala fallback"""
    name_norm = norm_name(name)
    for (c, mt, n), (p1, p2) in price_map.items():
        if mt == meal_type and norm_name(n) == name_norm:
            if c == city:
                return (p1, p2)
            if c == fallback_city and city == "shahdag":
                return (p1, p2)  # shahdag üçün qabala fiyatları
    return (None, None)


def main():
    price_map = extract_prices_from_excel()
    print("Excel-dən %d fiyat oxundu" % len(price_map))

    # meals.sql strukturu - mövcud INSERT-ləri price ilə yenilə
    with open(OUT_SQL, "r", encoding="utf-8") as f:
        content = f.read()
    # Əvvəlki işlədilmədən qalan qiymət blokunu at — əks halda hər işlədilmədə blok yenidən əlavə olunurdu
    # və DB-də hər yemək üçün təkrarlanan qiymət sətirləri yaranırdı (builder-də yeməklər 3 dəfə görünürdü).
    content = content.split(PRICES_MARKER)[0].rstrip() + "\n"

    # Hər ('city', 'meal_type', 'name', sort_order) - price əlavə et
    # Format: ('city', 'meal_type', 'name', 1) -> ('city', 'meal_type', 'name', 11, 11, 1)  price, price2, sort_order
    # Amma tour_meal_options-da yalnız price var; price2 tour_meal_prices-dadır.
    # Sadə: price əlavə et tour_meal_options-a. tour_meal_prices ayrıca INSERT ilə.

    def replacer(m):
        city = m.group(1)
        meal_type = m.group(2)
        name = m.group(3)
        sort_order = m.group(4)
        p1, p2 = get_price(price_map, city, meal_type, name)
        if p1 is not None:
            return "('%s', '%s', '%s', %.2f, %s)" % (city, meal_type, esc(name), p1, sort_order)
        return "('%s', '%s', '%s', NULL, %s)" % (city, meal_type, esc(name), sort_order)

    # Match both: (city, meal_type, name, sort_order) OR (city, meal_type, name, price_or_null, sort_order)
    new_content = re.sub(
        r"\('([^']+)', '([^']+)', '([^']*(?:''[^']*)*)', (?:NULL|\d+\.?\d*), (\d+)\)",
        replacer,
        content,
    )
    new_content = re.sub(
        r"\('([^']+)', '([^']+)', '([^']*(?:''[^']*)*)', (\d+)\)",
        replacer,
        new_content,
    )

    # INSERT sütun adlarını yenilə
    new_content = new_content.replace(
        "INSERT INTO `tour_meal_options` (`city`, `meal_type`, `name`, `sort_order`) VALUES",
        "INSERT INTO `tour_meal_options` (`city`, `meal_type`, `name`, `price`, `sort_order`) VALUES",
        # Hər yerdə
    )
    new_content = new_content.replace(
        "(`city`, `meal_type`, `name`, `sort_order`) VALUES",
        "(`city`, `meal_type`, `name`, `price`, `sort_order`) VALUES",
    )

    # Duplicate əvəzetmə (bəlkə bir neçə INSERT var)
    new_content = new_content.replace(
        "(`city`, `meal_type`, `name`, `sort_order`)",
        "(`city`, `meal_type`, `name`, `price`, `sort_order`)",
    )

    with open(OUT_SQL, "w", encoding="utf-8") as f:
        f.write(new_content)

    # tour_meal_prices üçün əlavə SQL - INSERT ... SELECT
    prices_sql_lines = [
        "",
        PRICES_MARKER,
        "INSERT INTO `tour_meal_prices` (`meal_option_id`, `season_start`, `season_end`, `price`, `price2`)",
    ]
    first = True
    for (city, meal_type, name), (p1, p2) in sorted(price_map.items(), key=lambda x: (x[0][0], x[0][1], x[0][2])):
        sel = "SELECT id, '2026-01-01', '2026-12-31', %.2f, %.2f FROM tour_meal_options WHERE city='%s' AND meal_type='%s' AND name='%s'" % (
            p1, p2, city, meal_type, esc(name).replace("'", "''")
        )
        if first:
            prices_sql_lines.append(sel)
            first = False
        else:
            prices_sql_lines.append("UNION ALL " + sel)

    if not first:
        prices_sql_lines[-1] = prices_sql_lines[-1] + " ON DUPLICATE KEY UPDATE `price` = VALUES(`price`), `price2` = VALUES(`price2`);"
        with open(OUT_SQL, "a", encoding="utf-8") as f:
            f.write("\n".join(prices_sql_lines))
            f.write(CITY_CODES_SQL)

    print("meals.sql yazıldı (price ilə)")


if __name__ == "__main__":
    main()
