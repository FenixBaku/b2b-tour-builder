# -*- coding: utf-8 -*-
"""Sembark Hotel Listing Excel-dan tour_hotels və tour_hotel_prices üçün SQL yaradır.
   Otellər siyahısı + hər otel/otaq/mövsüm üçün qiymət."""
import re
import openpyxl
import os
import sys

DEFAULT_EXCEL = os.path.join(os.path.dirname(__file__), "Sembark Hotel Listing _Revise 10.12.2025 (1).xlsx")
DEFAULT_DIR = os.path.join(os.path.dirname(__file__), "migrations")
EXCEL_PATH = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL
OUTPUT_DIR = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_DIR
OUTPUT_HOTELS = os.path.join(OUTPUT_DIR, "005_seed_hotels.sql")
OUTPUT_PRICES = os.path.join(OUTPUT_DIR, "012_seed_hotel_prices.sql")


def escape_sql(s):
    if s is None:
        return ""
    t = str(s).strip()
    return t.replace("\\", "\\\\").replace("'", "''").replace("\r", " ").replace("\n", " ")


def find_column(header_row, keywords):
    """Header sətirində açar sözlərə uyğun sütun indeksini tap"""
    for i, cell in enumerate(header_row):
        if cell:
            val = str(cell).strip().lower()
            for kw in keywords:
                if kw in val:
                    return i
    return None


def find_hotel_column(header_row):
    for i, cell in enumerate(header_row):
        if cell:
            val = str(cell).strip().lower()
            if val in ('name', 'hotel', 'otel'):
                return i
    for i, cell in enumerate(header_row):
        if cell:
            val = str(cell).strip().lower()
            if 'name' in val or 'hotel' in val:
                return i
    return None


def find_room_column(header_row):
    return find_column(header_row, ['room', 'otaq', 'room type'])


def find_price_columns(header_row):
    """Peak Season, Low Season və ya Price sütunlarını tap. [(col_idx, season_key), ...]"""
    out = []
    for i, cell in enumerate(header_row):
        if cell:
            val = str(cell).strip().lower()
            if 'peak' in val:
                out.append((i, 'peak'))
            elif 'low' in val or 'low season' in val:
                out.append((i, 'low'))
            elif re.match(r'^season\s*_?\s*2$', val):
                out.append((i, 'low'))
            elif val == 'price' or val == 'rate':
                out.append((i, 'standard'))
    return out if out else None


def sheet_to_region(sheet_name):
    s = str(sheet_name).strip().lower()
    if 'gabala' in s or 'qabala' in s:
        return 'gabala'
    if 'shahdag' in s or 'shahdag' in s:
        return 'shahdag'
    if 'sheki' in s or 'shaki' in s:
        return 'sheki'
    if 'shamakhi' in s or 'shamaxi' in s:
        return 'shamakhi'
    return 'baku'


def parse_price(val):
    if val is None:
        return None
    s = str(val).strip().replace(',', '.')
    try:
        return float(re.sub(r'[^\d.]', '', s))
    except (ValueError, TypeError):
        return None


def is_valid_hotel_name(name):
    if not name or len(name) < 2:
        return False
    n = name.strip()
    if len(n) > 120:
        return False
    lower = n.lower()
    if n.replace(" ", "").replace(".", "").isdigit():
        return False
    if lower in ('bb', 'hb', 'fb', 'ai', 'extras', 'room', 'suite'):
        return False
    if any(x in lower for x in ('2025', '2026', '01 jan', '01 june', '31 dec', ' - 31')):
        return False
    if re.search(r'room\s*\(\d+p\)', lower) or ' (2p)' in lower or ' (1p)' in lower:
        return False
    if not any(c.isalpha() for c in n):
        return False
    return True


def collect_hotels_and_prices(path):
    """Excel-dan otel adları və fiyatları topla"""
    hotels = set()
    prices = []  # [(hotel, region, room_type, season, price), ...]

    if not os.path.isfile(path):
        print(f"Fayl tapılmadı: {path}")
        return sorted(hotels), prices

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        region = sheet_to_region(sheet_name)
        header = rows[0]
        hotel_col = find_hotel_column(header)
        room_col = find_room_column(header)
        price_cols = find_price_columns(header)

        for row in rows[6:]:
            if not row or len(row) <= (hotel_col if hotel_col is not None else 0):
                continue

            hotel_val = row[hotel_col] if hotel_col is not None else str(sheet_name)
            hotel_str = str(hotel_val).strip() if hotel_val else ""

            if not hotel_str and hotel_col is None:
                hotel_str = str(sheet_name).strip()

            if hotel_str and is_valid_hotel_name(hotel_str):
                hotels.add(hotel_str)

                # Fiyatları topla
                room_type = ""
                if room_col is not None and room_col < len(row) and row[room_col]:
                    room_type = str(row[room_col]).strip()

                if price_cols:
                    for col_idx, season_key in price_cols:
                        if col_idx < len(row):
                            p = parse_price(row[col_idx])
                            if p is not None and p > 0:
                                prices.append((hotel_str, region, room_type or 'Standard', season_key, p))
                elif room_col is None and not price_cols:
                    # Yalnız otel adı var - fiyat sütunlarını header-dan yenidən yoxla
                    for i, cell in enumerate(header):
                        if cell and i < len(row):
                            val = str(cell).strip().lower()
                            if any(x in val for x in ['price', 'rate', 'peak', 'low']):
                                p = parse_price(row[i])
                                if p is not None and p > 0:
                                    prices.append((hotel_str, region, 'Standard', 'peak' if 'peak' in val else 'low', p))

    wb.close()
    return sorted(hotels), prices


def write_hotels_sql(hotels, sql_path):
    lines = [
        "-- Sembark Hotel Listing Excel-dan import",
        "SET NAMES utf8mb4;",
        ""
    ]
    if hotels:
        lines.append("INSERT IGNORE INTO `tour_hotels` (`name`, `region`, `sort_order`) VALUES")
        vals = []
        for i, name in enumerate(hotels, 1):
            region = 'baku'
            if ' - Gabala' in name or ' - gabala' in name:
                region = 'gabala'
            elif ' - Shahdag' in name or ' - shahdag' in name:
                region = 'shahdag'
            vals.append("('%s', '%s', %d)" % (escape_sql(name), escape_sql(region), i))
        lines.append(",\n".join(vals) + ";")
    else:
        lines.append("-- Heç bir otel tapılmadı.")
    os.makedirs(os.path.dirname(sql_path), exist_ok=True)
    with open(sql_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"SQL yazıldı: {sql_path} | {len(hotels)} otel")


def write_prices_sql(prices, sql_path):
    if not prices:
        return
    lines = [
        "-- Sembark Hotel Listing Excel-dan otaq fiyatları",
        "SET NAMES utf8mb4;",
        ""
    ]
    lines.append("INSERT INTO `tour_hotel_prices` (`hotel_name`, `region`, `room_type`, `season_type`, `price`) VALUES")
    vals = []
    for h, r, room, season, p in prices:
        vals.append("('%s', '%s', '%s', '%s', %.2f)" % (escape_sql(h), escape_sql(r), escape_sql(room), escape_sql(season), p))
    lines.append(",\n".join(vals))
    lines.append("ON DUPLICATE KEY UPDATE `price` = VALUES(`price`);")
    os.makedirs(os.path.dirname(sql_path), exist_ok=True)
    with open(sql_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"SQL yazıldı: {sql_path} | {len(prices)} fiyat")


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "--debug":
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))[:20]
            print(f"\n=== Sheet: {sheet_name} ===")
            for i, row in enumerate(rows):
                print(f"  {i}: {row[:14]}")
            if rows:
                header = rows[0]
                print("  --- Sütun axtarışı ---")
                print("  Hotel col:", find_hotel_column(header))
                print("  Room col:", find_room_column(header))
                print("  Price cols:", find_price_columns(header))
        wb.close()
        return

    hotels, prices = collect_hotels_and_prices(EXCEL_PATH)
    write_hotels_sql(hotels, OUTPUT_HOTELS)
    write_prices_sql(prices, OUTPUT_PRICES)


if __name__ == "__main__":
    main()
